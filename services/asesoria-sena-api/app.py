import io, os, re, sqlite3, unicodedata
from datetime import datetime, timezone
from pathlib import Path
from flask import Flask, request, jsonify, send_file, abort
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from werkzeug.utils import secure_filename

app=Flask(__name__)
DB=Path(os.environ.get('DB_PATH','/data/asesoria.sqlite3'))
UPLOAD_DIR=Path(os.environ.get('UPLOAD_DIR','/uploads'))
ADMIN_KEY=os.environ.get('ADMIN_KEY','')
ALLOWED_ORIGIN='https://jorgecotera.github.io'
MAX_UPLOAD=5*1024*1024

SCHEMA='''
CREATE TABLE IF NOT EXISTS solicitudes(
 id INTEGER PRIMARY KEY AUTOINCREMENT,
 code TEXT UNIQUE,
 nombre TEXT NOT NULL,
 familiar TEXT,
 telefono TEXT NOT NULL,
 correo TEXT,
 programa TEXT NOT NULL,
 fecha_atencion TEXT NOT NULL,
 documento_tipo TEXT,
 documento_numero TEXT,
 pago_declarado TEXT NOT NULL DEFAULT 'pendiente',
 pago_metodo TEXT,
 pago_referencia TEXT,
 pago_estado TEXT NOT NULL DEFAULT 'pendiente',
 pago_verificado_at TEXT,
 comprobante_path TEXT,
 estado TEXT NOT NULL DEFAULT 'solicitud_recibida',
 observaciones TEXT,
 consentimiento_at TEXT NOT NULL,
 created_at TEXT NOT NULL,
 updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_solicitudes_fecha ON solicitudes(fecha_atencion);
CREATE INDEX IF NOT EXISTS idx_solicitudes_pago ON solicitudes(pago_estado);
'''

def conn():
    c=sqlite3.connect(DB)
    c.row_factory=sqlite3.Row
    return c

def init_db():
    DB.parent.mkdir(parents=True,exist_ok=True); UPLOAD_DIR.mkdir(parents=True,exist_ok=True)
    with conn() as c: c.executescript(SCHEMA)
init_db()

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')

def cors(resp):
    origin=request.headers.get('Origin')
    if origin==ALLOWED_ORIGIN:
        resp.headers['Access-Control-Allow-Origin']=origin
        resp.headers['Vary']='Origin'
        resp.headers['Access-Control-Allow-Headers']='Content-Type, X-Admin-Key'
        resp.headers['Access-Control-Allow-Methods']='GET, POST, PATCH, OPTIONS'
    return resp
app.after_request(cors)

@app.route('/<path:path>',methods=['OPTIONS'])
@app.route('/',methods=['OPTIONS'])
def options(path=''): return ('',204)

def admin_required():
    if not ADMIN_KEY or request.headers.get('X-Admin-Key')!=ADMIN_KEY: abort(401)

def clean(v,limit=500): return (v or '').strip()[:limit]

def valid_date(v):
    try: datetime.strptime(v,'%Y-%m-%d'); return True
    except Exception: return False

@app.get('/health')
def health(): return jsonify(ok=True,service='asesoria-sena')

@app.post('/solicitudes')
def create_request():
    f=request.form
    nombre=clean(f.get('nombre'),180); telefono=clean(f.get('telefono'),60)
    programa=clean(f.get('programa'),240); fecha=clean(f.get('fecha'),10)
    if not nombre or not telefono or not programa or not valid_date(fecha):
        return jsonify(error='Faltan datos obligatorios o la fecha no es válida.'),400
    if f.get('acepta') not in ('on','true','1','si','sí'):
        return jsonify(error='Debe aceptar las condiciones del servicio.'),400
    declared=clean(f.get('pago'),60).lower()
    pago_estado='pendiente_verificacion' if 'realic' in declared or 'presencial' in declared else 'pendiente'
    receipt=request.files.get('comprobante')
    saved=None
    ts=now()
    with conn() as c:
        cur=c.execute('''INSERT INTO solicitudes(nombre,familiar,telefono,correo,programa,fecha_atencion,documento_tipo,documento_numero,pago_declarado,pago_metodo,pago_referencia,pago_estado,estado,observaciones,consentimiento_at,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',(
          nombre,clean(f.get('familiar'),180),telefono,clean(f.get('correo'),180),programa,fecha,
          clean(f.get('documento_tipo'),30),clean(f.get('documento_numero'),60),declared,
          clean(f.get('pago_metodo'),40),clean(f.get('referencia'),120),pago_estado,'solicitud_recibida',clean(f.get('observaciones'),1500),ts,ts,ts))
        rid=cur.lastrowid
        code=f"SOL-{datetime.now().strftime('%Y%m%d')}-{rid:04d}"
        c.execute('UPDATE solicitudes SET code=? WHERE id=?',(code,rid))
        if receipt and receipt.filename:
            receipt.stream.seek(0,2); size=receipt.stream.tell(); receipt.stream.seek(0)
            if size>MAX_UPLOAD: return jsonify(error='El comprobante supera el máximo de 5 MB.'),413
            ext=Path(secure_filename(receipt.filename)).suffix.lower()
            if ext not in ('.jpg','.jpeg','.png','.webp','.pdf'): return jsonify(error='Formato de comprobante no permitido.'),400
            saved=f'{code}{ext}'; receipt.save(UPLOAD_DIR/saved)
            c.execute('UPDATE solicitudes SET comprobante_path=? WHERE id=?',(saved,rid))
    return jsonify(ok=True,code=code,status='pendiente_verificacion' if pago_estado!='pendiente' else 'pendiente'),201

@app.get('/admin/solicitudes')
def list_requests():
    admin_required(); fecha=clean(request.args.get('fecha'),10); estado=clean(request.args.get('estado'),40)
    sql='SELECT * FROM solicitudes WHERE 1=1'; args=[]
    if fecha: sql+=' AND fecha_atencion=?'; args.append(fecha)
    if estado: sql+=' AND pago_estado=?'; args.append(estado)
    sql+=' ORDER BY fecha_atencion, CASE WHEN pago_verificado_at IS NULL THEN 1 ELSE 0 END, pago_verificado_at, created_at'
    with conn() as c: rows=[dict(r) for r in c.execute(sql,args)]
    for r in rows: r['has_comprobante']=bool(r.pop('comprobante_path',None))
    return jsonify(items=rows)

@app.patch('/admin/solicitudes/<int:rid>')
def update_request(rid):
    admin_required(); data=request.get_json(silent=True) or {}; allowed={}
    if 'pago_estado' in data and data['pago_estado'] in ('pendiente','pendiente_verificacion','confirmado','rechazado'):
        allowed['pago_estado']=data['pago_estado']; allowed['pago_verificado_at']=now() if data['pago_estado']=='confirmado' else None
        if data['pago_estado']=='confirmado': allowed['estado']='confirmado'
    if 'estado' in data and data['estado'] in ('solicitud_recibida','confirmado','atendido','cancelado'):
        allowed['estado']=data['estado']
    if 'observaciones' in data: allowed['observaciones']=clean(data['observaciones'],1500)
    if not allowed: return jsonify(error='Sin cambios válidos.'),400
    allowed['updated_at']=now(); setsql=', '.join(f'{k}=?' for k in allowed)
    with conn() as c:
        cur=c.execute(f'UPDATE solicitudes SET {setsql} WHERE id=?',list(allowed.values())+[rid])
        if not cur.rowcount: abort(404)
    return jsonify(ok=True)

@app.get('/admin/resumen')
def summary():
    admin_required(); fecha=clean(request.args.get('fecha'),10)
    where=' WHERE fecha_atencion=?' if fecha else ''; args=[fecha] if fecha else []
    with conn() as c:
        total=c.execute('SELECT COUNT(*) FROM solicitudes'+where,args).fetchone()[0]
        conf=c.execute("SELECT COUNT(*) FROM solicitudes"+where+(" AND" if where else " WHERE")+" pago_estado='confirmado'",args).fetchone()[0]
        pending=c.execute("SELECT COUNT(*) FROM solicitudes"+where+(" AND" if where else " WHERE")+" pago_estado IN ('pendiente','pendiente_verificacion')",args).fetchone()[0]
    return jsonify(total=total,confirmados=conf,pendientes=pending,valor_confirmado=conf*100000)

MONTHS=['','Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']
DAYS=['Lunes','Martes','Miércoles','Jueves','Viernes','Sábado','Domingo']

def export_rows(fecha=None,all_rows=False):
    sql="SELECT * FROM solicitudes WHERE 1=1"; args=[]
    if fecha: sql+=' AND fecha_atencion=?'; args.append(fecha)
    if not all_rows: sql+=" AND pago_estado='confirmado'"
    sql+=' ORDER BY COALESCE(pago_verificado_at,created_at), created_at'
    with conn() as c: return [dict(r) for r in c.execute(sql,args)]

def build_excel(rows,fecha=None):
    wb=Workbook(); ws=wb.active
    if fecha and valid_date(fecha):
        d=datetime.strptime(fecha,'%Y-%m-%d'); ws.title=f'{d.day} {MONTHS[d.month]} - {DAYS[d.weekday()]}'[:31]
    else: ws.title='Solicitudes'
    headers=['ORDEN','SOLICITUD','NOMBRE','FAMILIAR','TELÉFONO','EMAIL','PAGO','ESTADO PAGO','FECHA PAGO','FECHA ATENCIÓN','DOC.','N° DE DOCUMENTO','PROGRAMA','ESTADO','OBSERVACIONES','REGISTRO']
    yellow=PatternFill('solid',fgColor='FFFF00'); blue=PatternFill('solid',fgColor='0070C0'); green=PatternFill('solid',fgColor='92D050'); lightblue=PatternFill('solid',fgColor='BDD7EE'); orange=PatternFill('solid',fgColor='F4B183'); gray=PatternFill('solid',fgColor='E7E6E6')
    thin=Side(style='thin',color='000000'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
    for c,h in enumerate(headers,1):
        cell=ws.cell(1,c,h); cell.fill=yellow; cell.font=Font(bold=True); cell.alignment=Alignment(horizontal='center'); cell.border=border
    for i,r in enumerate(rows,1):
        vals=[i,r['code'],r['nombre'],r['familiar'],r['telefono'],r['correo'],r['pago_metodo'] or r['pago_declarado'],r['pago_estado'],r['pago_verificado_at'],r['fecha_atencion'],r['documento_tipo'],r['documento_numero'],r['programa'],r['estado'],r['observaciones'],r['created_at']]
        for c,v in enumerate(vals,1):
            cell=ws.cell(i+1,c,v); cell.border=border; cell.alignment=Alignment(vertical='top',wrap_text=c in (13,15))
            if c==1: cell.fill=blue; cell.font=Font(bold=True); cell.alignment=Alignment(horizontal='center')
            elif c==3: cell.fill=green
            elif c in (4,5,6): cell.fill=lightblue
            elif c in (7,8,9): cell.fill=yellow
            elif c in (10,11,12): cell.fill=lightblue
            elif c==13: cell.fill=orange
            elif c in (14,15,16): cell.fill=gray
    widths=[9,22,34,28,16,28,16,20,22,18,10,20,48,20,45,24]
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='D2'; ws.auto_filter.ref=ws.dimensions
    return wb

@app.get('/admin/export.xlsx')
def export_excel():
    admin_required(); fecha=clean(request.args.get('fecha'),10); all_rows=request.args.get('all')=='1'
    rows=export_rows(fecha or None,all_rows)
    wb=build_excel(rows,fecha or None); bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    if fecha and valid_date(fecha):
        d=datetime.strptime(fecha,'%Y-%m-%d'); filename=f'{d.day} de {MONTHS[d.month]} de {d.year}.xlsx'
    else: filename='Solicitudes SENA.xlsx'
    return send_file(bio,as_attachment=True,download_name=filename,mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
